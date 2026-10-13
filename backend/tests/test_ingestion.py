"""Unit tests for knowledge ingestion, conflict detection, speech normalization, and storage."""

from __future__ import annotations

import io
import json
import pytest

from app.schemas.ingestion import (
    ConflictType,
    FactItem,
    FactVerificationStatus,
)
from app.services.ingestion import (
    detect_conflicts,
    extract_text_from_file,
    parse_linkedin_export,
    review_fact,
    store_and_reindex_facts,
)
from app.services.storage import StorageAdapter
from app.services.tts import (
    AVAILABLE_VOICES,
    get_available_voices,
    prepare_text_for_speech,
)


def test_extract_text_txt() -> None:
    raw = b"Hello world! This is a test resume document."
    text = extract_text_from_file("resume.txt", raw)
    assert "Hello world!" in text


def test_extract_text_json() -> None:
    data = {"name": "Test Candidate", "skills": ["Python", "FastAPI"]}
    raw = json.dumps(data).encode("utf-8")
    text = extract_text_from_file("profile.json", raw)
    assert "Test Candidate" in text
    assert "Python" in text


def test_extract_text_docx() -> None:
    import docx
    doc = docx.Document()
    doc.add_paragraph("Engineering Experience at EchoMind")
    doc.add_paragraph("Developed RAG pipelines.")
    buf = io.BytesIO()
    doc.save(buf)
    raw = buf.getvalue()

    text = extract_text_from_file("test.docx", raw)
    assert "Engineering Experience at EchoMind" in text
    assert "Developed RAG pipelines." in text


def test_extract_text_pdf() -> None:
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    writer.write(buf)
    raw = buf.getvalue()

    # Should not throw exception on valid PDF
    text = extract_text_from_file("test.pdf", raw)
    assert isinstance(text, str)


def test_parse_linkedin_export() -> None:
    sample_linkedin = """
About
Final-year B.Tech student passionate about AI/ML and RAG systems.

Experience
AI/ML Engineer Intern at SkillDisk
Worked on phishing URL classification and machine-learning models.

Education
UVCE Bengaluru
Bachelor of Technology - BTech, Computer Science Engineering
2023 - 2027

Skills
Python
FastAPI
Machine Learning
"""
    facts = parse_linkedin_export(sample_linkedin, "linkedin_profile.pdf")
    assert len(facts) >= 3

    types = {f.entity_type for f in facts}
    assert "summary" in types
    assert "experience" in types
    assert "education" in types
    assert "skill" in types

    for f in facts:
        assert f.source == "linkedin_export"
        assert f.verification_status == FactVerificationStatus.PENDING_REVIEW


def test_conflict_detection_date_mismatch() -> None:
    # UVCE graduation is 2027 in verified profile; if imported fact claims 2024, it must flag a blocker
    conflicting_fact = FactItem(
        id="test-conflict-edu",
        entity_type="education",
        title="Graduated Early",
        content="Completed UVCE Computer Science degree in 2024.",
        source="resume_file",
        source_url_or_file="bad_resume.pdf",
        import_date="2026-10-13T00:00:00Z",
        verification_status=FactVerificationStatus.PENDING_REVIEW,
    )

    conflicts = detect_conflicts(conflicting_fact, [])
    assert any(c.conflict_type == ConflictType.DATE_MISMATCH for c in conflicts)
    assert any(c.severity == "blocker" for c in conflicts)


def test_conflict_detection_unsupported_claim() -> None:
    # Exaggerated revenue/metrics without verified proof must flag a warning
    unsupported_fact = FactItem(
        id="test-unsupported",
        entity_type="project",
        title="Unverified Startup",
        content="Built a crypto bot that generated $15 million in profits for 500k users.",
        source="unverified_upload",
        source_url_or_file="claim.txt",
        import_date="2026-10-13T00:00:00Z",
        verification_status=FactVerificationStatus.PENDING_REVIEW,
    )

    conflicts = detect_conflicts(unsupported_fact, [])
    assert any(c.conflict_type == ConflictType.UNSUPPORTED_CLAIM for c in conflicts)


def test_store_and_admin_review() -> None:
    new_fact = FactItem(
        id="test-review-item-1",
        entity_type="skill",
        title="Advanced PyTorch",
        content="Proficient in PyTorch and tensor operations.",
        source="manual",
        source_url_or_file="notes.txt",
        import_date="2026-10-13T00:00:00Z",
        verification_status=FactVerificationStatus.PENDING_REVIEW,
    )

    result = store_and_reindex_facts([new_fact])
    assert result.status == "success"

    # Admin approves
    approved = review_fact("test-review-item-1", "approve")
    assert approved is not None
    assert approved.verification_status == FactVerificationStatus.VERIFIED

    # Admin rejects
    rejected = review_fact("test-review-item-1", "reject")
    assert rejected is not None
    assert rejected.verification_status == FactVerificationStatus.REJECTED


def test_prepare_text_for_speech_normalizations() -> None:
    raw = (
        "I built a **FastAPI** backend with **RAG** and OCR for the AI/ML project! "
        "Verified Sources: 📌 [AI Learn](https://example.com) [1]"
    )
    speech = prepare_text_for_speech(raw)

    # Markdown stripped
    assert "**" not in speech
    assert "[" not in speech
    assert "]" not in speech
    assert "📌" not in speech
    assert "[1]" not in speech
    assert "Verified Sources:" not in speech

    # Pronunciation expanded
    assert "Fast A P I" in speech
    assert "rag" in speech
    assert "O C R" in speech
    assert "A I, M L" in speech


def test_available_voices_list() -> None:
    voices = get_available_voices()
    assert len(voices) >= 4
    voice_ids = {v["id"] for v in voices}
    assert "en-US-JennyNeural" in voice_ids
    assert "en-US-GuyNeural" in voice_ids


@pytest.mark.asyncio
async def test_storage_adapter_local_fallback() -> None:
    storage = StorageAdapter()
    session_id = "test-session-storage-123"

    turn = await storage.save_turn(
        session_id=session_id,
        role="user",
        content="Hello digital twin!",
    )
    assert turn["session_id"] == session_id
    assert turn["role"] == "user"

    turns = await storage.get_session_turns(session_id)
    assert len(turns) >= 1
    assert turns[0]["content"] == "Hello digital twin!"
